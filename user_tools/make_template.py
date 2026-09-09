#!/usr/bin/env python3
# ============================================================
# 生成 user_dataset_template/dataset_info.xlsx
#
# 三个工作表：
#   experiment      —— 实验信息（带下拉）
#   column_mapping  —— 列映射与单位（带下拉）
#   填写说明        —— 中文说明（只读参考，用户可不看）
#
# 运行：python -m user_tools.make_template
# ============================================================

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]

from user_tools.spec import (  # noqa: E402
    ALLOWED_CHEMISTRY,
    ALLOWED_CELL_CONFIGURATION,
    ALLOWED_COUNTER_ELECTRODE,
    ALLOWED_CURRENT_SIGN,
    ALLOWED_UNITS,
    ALLOWED_WORKING_ELECTRODE,
    CANONICAL_FIELDS,
    EXPERIMENT_FIELDS,
    FIELD_LABEL_ZH,
)

HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=11)
REQ_FILL = PatternFill("solid", fgColor="FFF2CC")
NOTE_FONT = Font(color="666666", size=10)


def _style_header(ws, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")


def build(path: Path) -> None:
    wb = Workbook()

    # ---------------------------------------------------------- experiment
    ws = wb.active
    ws.title = "experiment"
    ws.append(["field", "字段（中文）", "value", "填写说明"])
    for name, zh, required, allowed in EXPERIMENT_FIELDS:
        ws.append([name, zh, "", ""])
        r = ws.max_row
        ws.cell(row=r, column=1).font = Font(name="Consolas", size=10)
        note = "必填" if required else "可选"
        if allowed:
            note += "；只能填：" + " / ".join(allowed)
        ws.cell(row=r, column=4).value = note
        ws.cell(row=r, column=4).font = NOTE_FONT
        if required:
            ws.cell(row=r, column=3).fill = REQ_FILL
        if allowed:
            dv = DataValidation(
                type="list",
                formula1='"' + ",".join(allowed) + '"',
                allow_blank=True,
            )
            ws.add_data_validation(dv)
            dv.add(ws.cell(row=r, column=3))
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 30
    ws.column_dimensions["D"].width = 52
    _style_header(ws, 4)
    ws.freeze_panes = "A2"

    # ------------------------------------------------------ column_mapping
    ws2 = wb.create_sheet("column_mapping")
    ws2.append(["canonical_field", "中文", "source_column", "source_unit", "说明"])
    for f in CANONICAL_FIELDS:
        ws2.append([f, FIELD_LABEL_ZH[f], "", "", ""])
        r = ws2.max_row
        ws2.cell(row=r, column=1).font = Font(name="Consolas", size=10)
        note = "必填" if f in ("time", "current", "voltage") else "可选"
        ws2.cell(row=r, column=5).value = (
            f"{note}；单位只能填：" + " / ".join(ALLOWED_UNITS[f])
        )
        ws2.cell(row=r, column=5).font = NOTE_FONT
        if f in ("time", "current", "voltage"):
            ws2.cell(row=r, column=3).fill = REQ_FILL
            ws2.cell(row=r, column=4).fill = REQ_FILL
        dv = DataValidation(
            type="list",
            formula1='"' + ",".join(ALLOWED_UNITS[f]) + '"',
            allow_blank=True,
        )
        ws2.add_data_validation(dv)
        dv.add(ws2.cell(row=r, column=4))
    ws2.column_dimensions["A"].width = 18
    ws2.column_dimensions["B"].width = 12
    ws2.column_dimensions["C"].width = 28
    ws2.column_dimensions["D"].width = 14
    ws2.column_dimensions["E"].width = 46
    _style_header(ws2, 5)
    ws2.freeze_panes = "A2"

    # ------------------------------------------------------------ 填写说明
    ws3 = wb.create_sheet("填写说明")
    lines = [
        ("第一步：复制整个模板文件夹，改成你自己的名字（例如 我的电池数据）。", True),
        ("", False),
        ("第二步：把你的数据文件放进 raw 文件夹（.csv 或 .xls 或 .xlsx，只放一个）。", True),
        ("", False),
        ("第三步：填 experiment 表（黄色格子是必填）。", True),
        ("   chemistry        填 LFP / NMC / LCO / other", False),
        ("   cell_configuration   full_cell = 全电池；half_cell = 半电池", False),
        ("   working_electrode    半电池才填：你的被测电极是 positive 还是 negative", False),
        ("   counter_electrode    半电池才填：对电极（锂金属填 lithium_metal）", False),
        ("   current_sign     看你的原始数据：放电那一段电流是正数还是负数", False),
        ("                    discharge_positive = 放电记为正", False),
        ("                    discharge_negative = 放电记为负（多数设备是这种）", False),
        ("   voltage_lower / voltage_upper   你实际用的电压窗口（V）", False),
        ("   active_material_loading   半电池一定要填（面容量 mAh/cm2）", False),
        ("", False),
        ("第四步：填 column_mapping 表。", True),
        ("   source_column = 你数据文件里的列名（要一模一样，含空格和方括号）", False),
        ("   source_unit   = 那一列的单位（从下拉里选）", False),
        ("   time / current / voltage 三行必填，其余没有就留空", False),
        ("", False),
        ("第五步：双击「导入并检查数据.bat」，看 validation_report.html。", True),
        ("", False),
        ("第六步：报告说可以仿真时，再双击「运行仿真.bat」。", True),
        ("", False),
        ("注意：", True),
        ("   本工具不做拟合，不调参数。跑出来的 RMSE 只是差异大小的参考，", False),
        ("   不等于模型验证，也不等于找到了误差原因。", False),
    ]
    for text, bold in lines:
        ws3.append([text])
        if bold:
            ws3.cell(row=ws3.max_row, column=1).font = Font(bold=True, size=11)
    ws3.column_dimensions["A"].width = 96

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> int:
    pkg = ROOT / "user_dataset_template"
    (pkg / "raw").mkdir(parents=True, exist_ok=True)
    out = pkg / "dataset_info.xlsx"
    build(out)
    print(f"模板已生成：{out.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
